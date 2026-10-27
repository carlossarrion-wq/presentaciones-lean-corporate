#!/usr/bin/env python3
"""Prueba de humo del skill: verifica dependencias y que todos los scripts funcionan.

    python tests/smoke_test.py          # desde la carpeta del skill

Sale con código 0 si todo pasa. Las pruebas que requieren LibreOffice se marcan SKIP si no
está instalado (el skill funciona, pero sin recálculo/conversión).
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
S = ROOT / "scripts"
sys.path.insert(0, str(S))

results = []


def check(name, fn):
    try:
        r = fn()
        results.append((name, "SKIP" if r == "skip" else "OK", ""))
    except Exception as e:  # noqa: BLE001
        results.append((name, "FAIL", f"{type(e).__name__}: {e}"))


def run(*args):
    p = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)
    if p.returncode != 0:
        raise RuntimeError(p.stderr or p.stdout)
    return p.stdout


def main():
    tmp = Path(tempfile.mkdtemp(prefix="excel_skill_test_"))
    import _soffice
    has_lo = _soffice.find_soffice() is not None

    def deps():
        import openpyxl, pandas, lxml  # noqa: F401
    check("dependencias Python (openpyxl, pandas, lxml)", deps)
    check("LibreOffice disponible", lambda: None if has_lo else "skip")

    demo = tmp / "demo.xlsx"

    def build():
        run(S / "build_from_spec.py", ROOT / "assets" / "spec_example.json", demo, "--no-recalc")
        assert demo.exists()
    check("build_from_spec.py", build)

    def recalc():
        if not has_lo:
            return "skip"
        out = json.loads(run(S / "recalc.py", demo))
        assert out["status"] == "success" and out["total_errors"] == 0, out
        import openpyxl
        v = openpyxl.load_workbook(demo, data_only=True)["Ventas"]["D10"].value
        assert abs(v - 109081.5) < 1e-6, v
    check("recalc.py (valores correctos)", recalc)

    def recalc_errors():
        if not has_lo:
            return "skip"
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws["A1"] = 1; ws["A2"] = "=A1/0"; ws["A3"] = "=TEXTJOIN(\",\",TRUE,A1)"
        f = tmp / "err.xlsx"; wb.save(f)
        out = json.loads(run(S / "recalc.py", f, "--check-only"))
        assert out["status"] == "errors_found" and out["total_errors"] == 2, out
        assert any("TEXTJOIN" in w for w in out["warnings"])
    check("recalc.py (detección de errores y avisos)", recalc_errors)

    check("inspect_xlsx.py", lambda: json.loads(run(S / "inspect_xlsx.py", demo, "--json")))

    def audit():
        out = json.loads(run(S / "audit_formulas.py", demo, "--json"))
        assert out["by_type"].get("inconsistent", 0) == 0, out
    check("audit_formulas.py", audit)

    def patch_and_diff():
        p = tmp / "patched.xlsx"
        run(S / "patch_cells.py", demo, "-o", p, "--set", "Ventas!B5=20000", "--set", "Ventas!B12==D10*2")
        out = json.loads(run(S / "diff_xlsx.py", demo, p, "--json"))
        assert out["total_changes"] == 2, out
    check("patch_cells.py + diff_xlsx.py", patch_and_diff)

    def insert_safe():
        import openpyxl
        from xlsx_helpers import insert_rows_safe
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Hoja 1"; o = wb.create_sheet("R")
        for i in range(1, 6):
            ws.cell(i, 1, i)
        ws["A6"] = "=SUM(A1:A5)"; o["A1"] = "='Hoja 1'!A6"
        insert_rows_safe(wb, ws, 3, 2)
        assert ws["A8"].value == "=SUM(A1:A7)", ws["A8"].value
        assert o["A1"].value == "='Hoja 1'!A8", o["A1"].value
    check("xlsx_helpers.insert_rows_safe", insert_safe)

    def csv_conv():
        c = tmp / "es.csv"
        c.write_bytes("Fecha;Importe\n15/01/2026;1.234,56\n".encode("cp1252"))
        x = tmp / "es.xlsx"
        run(S / "convert.py", c, x)
        import openpyxl
        ws = openpyxl.load_workbook(x).active
        assert abs(ws["B2"].value - 1234.56) < 1e-9 and hasattr(ws["A2"].value, "year")
    check("convert.py CSV español → xlsx", csv_conv)

    def pdf():
        if not has_lo:
            return "skip"
        run(S / "convert.py", demo, tmp / "demo.pdf")
        assert (tmp / "demo.pdf").stat().st_size > 1000
    check("convert.py xlsx → pdf", pdf)

    width = max(len(n) for n, _, _ in results)
    for n, st, msg in results:
        print(f"[{st:^4}] {n:<{width}} {msg}")
    fails = [r for r in results if r[1] == "FAIL"]
    print(f"\n{len(results) - len(fails)}/{len(results)} correctas" + (f" · temporales en {tmp}" if fails else ""))
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
