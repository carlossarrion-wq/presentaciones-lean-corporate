#!/usr/bin/env python3
"""Genera un libro Excel profesional a partir de una especificación JSON declarativa.

Pensado para que un agente cree libros consistentes sin reescribir código de estilos.
Para lógica compleja (modelos con muchas fórmulas) escribe tu propio script con openpyxl
y xlsx_helpers; este generador cubre el 80 % de casos: tablas de datos, totales, supuestos,
validaciones, formato condicional y gráficos.

Uso:
    python build_from_spec.py spec.json salida.xlsx [--no-recalc]

Ver assets/spec_example.json para un ejemplo completo. Esquema resumido:
{
  "font": "Arial",                       # opcional
  "sheets": [{
    "name": "Ventas",
    "title": "Informe de ventas 2026",   # opcional: fila de título encima de la tabla
    "start_cell": "A1",                   # esquina de la tabla (se desplaza si hay título)
    "columns": [{"header": "Mes", "format": "month", "width": 12}, ...],
    "rows": [["2026-01-01", 1200, "=B{r}*Supuestos!$B$2"], ...],   # {r} = nº de fila actual
    "table": {"name": "TVentas", "style": "TableStyleMedium2"},     # o "header_style": true
    "totals": {"label": "Total", "functions": {"Importe": "SUM", "Margen": "AVERAGE"}},
    "freeze": true,
    "cells": {"H2": {"value": 0.21, "format": "pct", "input": true, "note": "Fuente: BOE"}},
    "conditional": [{"range": "C2:C13", "type": "color_scale" | "data_bars" |
                     "less_than"|"greater_than", "value": 0, "kind": "bad"}],
    "validations": [{"range": "D2:D200", "list": ["Sí", "No"]}],
    "charts": [{"type": "bar"|"line"|"pie", "data": "B1:B13", "categories": "A2:A13",
                "anchor": "F2", "title": "Ventas"}],
    "print": {"landscape": true}
  }]
}
Las cadenas ISO 'AAAA-MM-DD' en columnas con formato de fecha se convierten en fechas reales.
Al terminar se ejecuta recalc.py (si hay fórmulas y LibreOffice está disponible).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string  # noqa: E402

import xlsx_helpers as H  # noqa: E402

DATE_FORMATS = {"date", "date_iso", "datetime", "month"}


def coerce(value, fmt):
    if isinstance(value, str) and fmt in DATE_FORMATS and re.fullmatch(r"\d{4}-\d{2}-\d{2}(T.*)?", value):
        return dt.datetime.fromisoformat(value)
    return value


def build_sheet(wb, spec, idx):
    ws = wb.active if idx == 0 else wb.create_sheet()
    ws.title = spec["name"][:31]
    col_letter, row0 = coordinate_from_string(spec.get("start_cell", "A1"))
    col0 = column_index_from_string(col_letter)
    if spec.get("title"):
        c = ws.cell(row0, col0, spec["title"])
        c.font = H.font(bold=True, size=14)
        row0 += 2
    columns = spec.get("columns", [])
    rows = spec.get("rows", [])
    has_formulas = False
    if columns:
        for j, colspec in enumerate(columns):
            ws.cell(row0, col0 + j, colspec["header"])
        for i, row in enumerate(rows, start=1):
            r = row0 + i
            for j, v in enumerate(row):
                fmt = columns[j].get("format") if j < len(columns) else None
                if isinstance(v, str) and v.startswith("="):
                    v = v.replace("{r}", str(r))
                    has_formulas = True
                cell = ws.cell(r, col0 + j, coerce(v, fmt))
                if fmt:
                    cell.number_format = H.NUMBER_FORMATS.get(fmt, fmt)
                if columns[j].get("input") if j < len(columns) else False:
                    H.mark_input(cell)
        last_row = row0 + len(rows)
        last_col = col0 + len(columns) - 1
        ref = f"{get_column_letter(col0)}{row0}:{get_column_letter(last_col)}{max(last_row, row0 + 1)}"
        table = spec.get("table")
        if table and rows:
            H.add_table(ws, ref, table.get("name", f"T_{ws.title}"), table.get("style", "TableStyleMedium2"))
        # Estilo explícito también con tabla: LibreOffice/Google Sheets no pintan estilos de tabla
        H.style_header_row(ws, row0, col0, last_col)
        totals = spec.get("totals")
        if totals and rows:
            tr = last_row + 1
            label_cell = ws.cell(tr, col0, totals.get("label", "Total"))
            label_cell.font = H.font(bold=True)
            headers = [c["header"] for c in columns]
            for name, func in totals.get("functions", {}).items():
                j = headers.index(name)
                L = get_column_letter(col0 + j)
                c = ws.cell(tr, col0 + j, f"={func.upper()}({L}{row0 + 1}:{L}{last_row})")
                fmt = columns[j].get("format")
                if fmt:
                    c.number_format = H.NUMBER_FORMATS.get(fmt, fmt)
                c.font = H.font(bold=True)
                has_formulas = True
            for j in range(len(columns)):
                ws.cell(tr, col0 + j).border = H.Border(top=H.Side(style="thin"), bottom=H.Side(style="double"))
        if spec.get("freeze", True) and rows:
            ws.freeze_panes = ws.cell(row0 + 1, col0)
        H.autofit_columns(ws)
        for j, colspec in enumerate(columns):
            if colspec.get("width"):
                ws.column_dimensions[get_column_letter(col0 + j)].width = colspec["width"]

    for coord, cs in spec.get("cells", {}).items():
        if not isinstance(cs, dict):
            cs = {"value": cs}
        v = cs.get("value")
        if isinstance(v, str) and v.startswith("="):
            has_formulas = True
        c = ws[coord]
        c.value = coerce(v, cs.get("format"))
        if cs.get("format"):
            c.number_format = H.NUMBER_FORMATS.get(cs["format"], cs["format"])
        if cs.get("bold"):
            c.font = H.font(bold=True)
        if cs.get("input"):
            H.mark_input(c, key_assumption=cs.get("key", False))
        if cs.get("note"):
            H.add_comment(c, cs["note"])
    for cf in spec.get("conditional", []):
        t = cf["type"]
        if t == "color_scale":
            H.add_color_scale(ws, cf["range"])
        elif t == "data_bars":
            H.add_data_bars(ws, cf["range"])
        elif t in ("less_than", "greater_than", "equal"):
            op = {"less_than": "lessThan", "greater_than": "greaterThan", "equal": "equal"}[t]
            H.highlight_cells(ws, cf["range"], op, [cf["value"]], cf.get("kind", "bad"))
        elif t == "formula":
            H.highlight_formula(ws, cf["range"], cf["formula"], cf.get("kind", "neutral"))
    for v in spec.get("validations", []):
        if "list" in v or "source" in v:
            H.add_list_validation(ws, v["range"], v.get("list"), v.get("source"))
        else:
            H.add_number_validation(ws, v["range"], v.get("min"), v.get("max"), v.get("whole", False))
    for ch in spec.get("charts", []):
        kind = ch.get("type", "bar")
        kwargs = dict(title=ch.get("title", ""))
        if kind == "bar":
            H.add_bar_chart(ws, ch["data"], ch["categories"], ch["anchor"], horizontal=ch.get("horizontal", False),
                            stacked=ch.get("stacked", False), y_title=ch.get("y_title"), **kwargs)
        elif kind == "line":
            H.add_line_chart(ws, ch["data"], ch["categories"], ch["anchor"], y_title=ch.get("y_title"), **kwargs)
        elif kind == "pie":
            H.add_pie_chart(ws, ch["data"], ch["categories"], ch["anchor"], **kwargs)
    if spec.get("print"):
        H.set_print_setup(ws, landscape=spec["print"].get("landscape", True))
    return has_formulas


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("output")
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args()
    spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    wb = Workbook()
    H.apply_default_font(wb, spec.get("font", "Arial"))
    any_formulas = False
    for i, s in enumerate(spec["sheets"]):
        any_formulas |= build_sheet(wb, s, i)
    wb.save(a.output)
    print(f"Creado {a.output}")
    if any_formulas and not a.no_recalc:
        res = subprocess.run([sys.executable, str(HERE / "recalc.py"), a.output], capture_output=True, text=True)
        print(res.stdout or res.stderr)


if __name__ == "__main__":
    main()
