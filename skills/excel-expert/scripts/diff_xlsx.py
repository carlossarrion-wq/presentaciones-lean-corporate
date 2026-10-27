#!/usr/bin/env python3
"""Compara dos libros celda a celda. Úsalo tras editar para demostrar que SOLO cambiaron
las celdas previstas (y que no se rompieron fórmulas ajenas).

Uso:
    python diff_xlsx.py original.xlsx modificado.xlsx            # compara fórmulas/valores escritos
    python diff_xlsx.py original.xlsx modificado.xlsx --values   # compara valores calculados (cache)
    python diff_xlsx.py a.xlsx b.xlsx --sheet Ventas --max 500 --json

Requisito para --values: ambos ficheros deben tener valores cacheados (recalc.py).
"""
from __future__ import annotations

import argparse
import json
import warnings
import math

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def norm(v):
    if isinstance(v, float) and math.isfinite(v):
        return round(v, 9)
    if v is not None and not isinstance(v, (str, int, float, bool)):
        return getattr(v, "text", None) or str(v)
    return v


def diff(a_path, b_path, values=False, only_sheet=None, max_items=200):
    wa = openpyxl.load_workbook(a_path, data_only=values)
    wb = openpyxl.load_workbook(b_path, data_only=values)
    res = {"sheets_added": [s for s in wb.sheetnames if s not in wa.sheetnames],
           "sheets_removed": [s for s in wa.sheetnames if s not in wb.sheetnames],
           "changes": [], "total_changes": 0}
    for name in wa.sheetnames:
        if name not in wb.sheetnames or (only_sheet and name != only_sheet):
            continue
        A, B = wa[name], wb[name]
        max_r = max(A.max_row, B.max_row)
        max_c = max(A.max_column, B.max_column)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                va, vb = norm(A.cell(r, c).value), norm(B.cell(r, c).value)
                if va != vb:
                    res["total_changes"] += 1
                    if len(res["changes"]) < max_items:
                        res["changes"].append({"sheet": name, "cell": B.cell(r, c).coordinate,
                                               "before": va, "after": vb})
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("a")
    ap.add_argument("b")
    ap.add_argument("--values", action="store_true")
    ap.add_argument("--sheet")
    ap.add_argument("--max", type=int, default=200)
    ap.add_argument("--json", action="store_true")
    x = ap.parse_args()
    res = diff(x.a, x.b, x.values, x.sheet, x.max)
    if x.json:
        print(json.dumps(res, ensure_ascii=False, indent=2, default=str))
        return
    if res["sheets_added"]:
        print("Hojas añadidas:", res["sheets_added"])
    if res["sheets_removed"]:
        print("Hojas eliminadas:", res["sheets_removed"])
    print(f"Celdas distintas: {res['total_changes']}")
    for ch in res["changes"]:
        print(f"  {ch['sheet']}!{ch['cell']}: {ch['before']!r} → {ch['after']!r}")
    if res["total_changes"] > len(res["changes"]):
        print(f"  … y {res['total_changes'] - len(res['changes'])} más (usa --max)")


if __name__ == "__main__":
    main()
