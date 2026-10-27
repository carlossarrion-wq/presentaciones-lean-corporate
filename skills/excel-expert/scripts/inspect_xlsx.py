#!/usr/bin/env python3
"""Radiografía de un libro Excel ANTES de editarlo o analizarlo.

Muestra, por hoja: rango usado, celdas con datos/fórmulas, celdas combinadas, paneles
inmovilizados, tablas, validaciones, formato condicional, gráficos, imágenes, filtros,
visibilidad y una vista previa de las primeras filas CON COORDENADAS. A nivel de libro:
nombres definidos, vínculos externos, macros y estilos de entrada (fuente/relleno) que
sugieren qué celdas son "inputs".

Uso:
    python inspect_xlsx.py libro.xlsx                    # resumen de todas las hojas
    python inspect_xlsx.py libro.xlsx --sheet Ventas     # solo una hoja
    python inspect_xlsx.py libro.xlsx --rows 30 --cols 15
    python inspect_xlsx.py libro.xlsx --formulas         # lista todas las fórmulas
    python inspect_xlsx.py libro.xlsx --json             # salida JSON (para encadenar)

Soporta .xlsx/.xlsm/.xltx/.xltm. Para .xls/.ods/.csv conviértelo antes (convert.py).
"""
from __future__ import annotations

import argparse
import json
import warnings
import sys
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def _fmt(v, width=18):
    if v is None:
        return ""
    if isinstance(v, float):
        v = round(v, 4)
    elif hasattr(v, "strftime"):
        v = v.strftime("%Y-%m-%d") if not (getattr(v, "hour", 0) or getattr(v, "minute", 0)) else v.strftime("%Y-%m-%d %H:%M")
    s = str(v).replace("\n", "⏎")
    return s if len(s) <= width else s[: width - 1] + "…"


def _style_signature(cell):
    f = cell.font
    color = None
    if f is not None and f.color is not None and f.color.type == "rgb":
        color = f.color.rgb
    fill = None
    if cell.fill is not None and cell.fill.fill_type == "solid" and cell.fill.fgColor.type == "rgb":
        fill = cell.fill.fgColor.rgb
    return color, fill


def inspect(path: Path, only_sheet: str | None, n_rows: int, n_cols: int, list_formulas: bool):
    keep_vba = path.suffix.lower() in (".xlsm", ".xltm")
    wb_f = openpyxl.load_workbook(path, keep_vba=keep_vba)           # fórmulas
    wb_v = openpyxl.load_workbook(path, data_only=True)              # valores cacheados

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    info = {
        "file": str(path),
        "sheets": wb_f.sheetnames,
        "active_sheet": wb_f.active.title if wb_f.active else None,
        "has_vba": any(n.endswith("vbaProject.bin") for n in names),
        "has_pivot_tables": any("pivotTable" in n for n in names),
        "has_slicers": any("slicer" in n.lower() for n in names),
        "external_links": len(getattr(wb_f, "_external_links", []) or []),
        "defined_names": {},
        "sheet_details": [],
    }
    try:
        for name, dn in wb_f.defined_names.items():
            info["defined_names"][name] = dn.attr_text
    except Exception:
        pass

    for ws in wb_f.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        wv = wb_v[ws.title]
        n_values = n_formulas = n_cached_missing = 0
        font_colors, fills = Counter(), Counter()
        formulas = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                is_formula = (isinstance(c.value, str) and c.value.startswith("=")) or \
                    type(c.value).__name__ in ("ArrayFormula", "DataTableFormula")
                if is_formula:
                    n_formulas += 1
                    cached = wv[c.coordinate].value
                    if cached is None:
                        n_cached_missing += 1
                    if list_formulas:
                        text = c.value if isinstance(c.value, str) else getattr(c.value, "text", str(c.value))
                        formulas.append({"cell": c.coordinate, "formula": text, "cached": cached})
                else:
                    n_values += 1
                    col, fill = _style_signature(c)
                    if col:
                        font_colors[col] += 1
                    if fill:
                        fills[fill] += 1

        preview = []
        max_r = min(ws.max_row, n_rows)
        max_c = min(ws.max_column, n_cols)
        for r in range(1, max_r + 1):
            cells = []
            for col in range(1, max_c + 1):
                cf = ws.cell(r, col).value
                cv = wv.cell(r, col).value
                if isinstance(cf, str) and cf.startswith("="):
                    cells.append(f"{_fmt(cv, 12)}{{={_fmt(cf[1:], 14)}}}" if cv is not None else f"{{{_fmt(cf, 22)}}}")
                else:
                    cells.append(_fmt(cv))
            preview.append(cells)

        detail = {
            "name": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "value_cells": n_values,
            "formula_cells": n_formulas,
            "formulas_without_cached_value": n_cached_missing,
            "merged_ranges": [str(m) for m in ws.merged_cells.ranges][:50],
            "freeze_panes": ws.freeze_panes,
            "auto_filter": ws.auto_filter.ref,
            "tables": {t.name: t.ref for t in ws.tables.values()},
            "data_validations": [
                {"sqref": str(dv.sqref), "type": dv.type, "formula1": dv.formula1}
                for dv in ws.data_validations.dataValidation
            ][:30],
            "conditional_formatting_ranges": [str(cf.sqref) for cf in ws.conditional_formatting][:30],
            "charts": len(getattr(ws, "_charts", [])),
            "images": len(getattr(ws, "_images", [])),
            "protected": bool(ws.protection.sheet),
            "column_widths": {k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}
            if ws.column_dimensions else {},
            "top_font_colors_in_values": font_colors.most_common(5),
            "top_fills_in_values": fills.most_common(5),
            "preview_header": [get_column_letter(c) for c in range(1, max_c + 1)],
            "preview": preview,
        }
        if list_formulas:
            detail["formulas"] = formulas
        info["sheet_details"].append(detail)
    return info


def print_human(info):
    print(f"# {info['file']}")
    print(f"Hojas: {info['sheets']}  | activa: {info['active_sheet']}")
    flags = []
    if info["has_vba"]:
        flags.append("MACROS VBA (abrir con keep_vba=True y guardar como .xlsm)")
    if info["has_pivot_tables"]:
        flags.append("TABLAS DINÁMICAS (openpyxl las conserva pero no las actualiza)")
    if info["has_slicers"]:
        flags.append("SEGMENTACIONES (openpyxl las PIERDE al guardar → usar patch_cells.py)")
    if info["external_links"]:
        flags.append(f"{info['external_links']} VÍNCULO(S) EXTERNO(S)")
    if flags:
        print("⚠ " + " | ".join(flags))
    if info["defined_names"]:
        print("Nombres definidos:", info["defined_names"])
    for d in info["sheet_details"]:
        print(f"\n## {d['name']}  [{d['state']}]  rango={d['dimensions']}  "
              f"valores={d['value_cells']} fórmulas={d['formula_cells']}"
              + (f" (sin valor cacheado: {d['formulas_without_cached_value']})"
                 if d['formulas_without_cached_value'] else ""))
        for key, label in [("merged_ranges", "Combinadas"), ("tables", "Tablas"),
                           ("data_validations", "Validaciones"),
                           ("conditional_formatting_ranges", "Formato condicional"),
                           ("column_widths", "Anchos")]:
            if d[key]:
                print(f"  {label}: {d[key]}")
        extras = []
        if d["freeze_panes"]:
            extras.append(f"inmovilizado={d['freeze_panes']}")
        if d["auto_filter"]:
            extras.append(f"autofiltro={d['auto_filter']}")
        if d["charts"]:
            extras.append(f"gráficos={d['charts']}")
        if d["images"]:
            extras.append(f"imágenes={d['images']}")
        if d["protected"]:
            extras.append("PROTEGIDA")
        if extras:
            print("  " + "  ".join(extras))
        if d["top_font_colors_in_values"] or d["top_fills_in_values"]:
            print(f"  Colores de fuente en valores: {d['top_font_colors_in_values']}  "
                  f"Rellenos: {d['top_fills_in_values']}  (pistas de celdas de entrada)")
        print("  Vista previa (valor{=fórmula}):")
        print("      | " + " | ".join(f"{h:<18}" for h in d["preview_header"]))
        for i, row in enumerate(d["preview"], 1):
            print(f"  {i:>3} | " + " | ".join(f"{v:<18}" for v in row))
        if "formulas" in d:
            print("  Fórmulas:")
            for f in d["formulas"]:
                print(f"    {f['cell']}: {f['formula']}   → {f['cached']}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--rows", type=int, default=15)
    ap.add_argument("--cols", type=int, default=10)
    ap.add_argument("--formulas", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    path = Path(a.file)
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        sys.exit(f"Formato {path.suffix} no soportado por openpyxl; usa convert.py a .xlsx primero.")
    info = inspect(path, a.sheet, a.rows, a.cols, a.formulas)
    if a.json:
        print(json.dumps(info, ensure_ascii=False, indent=2, default=str))
    else:
        print_human(info)


if __name__ == "__main__":
    main()
