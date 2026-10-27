#!/usr/bin/env python3
"""Conversión entre formatos de hoja de cálculo.

    python convert.py entrada.xls  salida.xlsx      # .xls/.ods/.xlsb/.xlsm → .xlsx (LibreOffice)
    python convert.py datos.csv    salida.xlsx      # CSV/TSV → .xlsx (detecta separador y codificación)
    python convert.py datos.csv    salida.xlsx --sep ";" --decimal "," --encoding latin-1
    python convert.py libro.xlsx   salida.csv  --sheet Ventas   # una hoja → CSV (valores)
    python convert.py libro.xlsx   carpeta/    --to csv         # todas las hojas → un CSV por hoja
    python convert.py libro.xlsx   salida.pdf                   # PDF con LibreOffice
    python convert.py libro.xlsx   salida.ods

Notas
- Exportar a CSV usa los valores CACHEADOS: si el libro lo generó openpyxl, ejecuta antes
  recalc.py o las celdas con fórmula saldrán vacías.
- CSV de salida en UTF-8 con BOM (utf-8-sig) para que Excel lo abra con tildes correctas;
  usa --sep ";" si el destinatario usa Excel en configuración regional española.
- .xlsx → .xls no se recomienda (límite 65.536 filas, se pierden funciones); se hace vía LibreOffice.
"""
from __future__ import annotations

import argparse
import csv
import shutil
import sys
import tempfile
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _soffice  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning)

LO_TARGETS = {".pdf": "pdf", ".ods": "ods", ".xls": "xls:MS Excel 97",
              ".xlsx": "xlsx:Calc Office Open XML"}


def sniff_csv(path: Path, encoding: str | None):
    raw = path.read_bytes()[:200_000]
    encs = [encoding] if encoding else ["utf-8-sig", "utf-8", "cp1252", "latin-1"]
    for enc in encs:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise RuntimeError("No se pudo detectar la codificación; usa --encoding")
    try:
        dialect = csv.Sniffer().sniff(text[:50_000], delimiters=",;\t|")
        sep = dialect.delimiter
    except csv.Error:
        sep = ","
    return enc, sep


def csv_to_xlsx(src: Path, dst: Path, sep, decimal, encoding, sheet_name):
    import pandas as pd
    from openpyxl import Workbook
    from xlsx_helpers import write_dataframe, apply_default_font

    enc, guessed_sep = sniff_csv(src, encoding)
    sep = sep or guessed_sep
    if decimal is None:
        decimal = "," if sep == ";" else "."
    thousands = "." if decimal == "," else None
    df = pd.read_csv(src, sep=sep, decimal=decimal, thousands=thousands, encoding=enc,
                     low_memory=False)
    # Intentar convertir columnas de texto que parezcan fechas
    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            sample = df[col].dropna().astype(str).head(50)
            if len(sample) and sample.str.match(r"^\d{1,4}[-/]\d{1,2}[-/]\d{1,4}").mean() > 0.9:
                parsed = pd.to_datetime(df[col], errors="coerce", dayfirst=not sample.str.match(r"^\d{4}").all())
                if parsed.notna().mean() > 0.9:
                    df[col] = parsed
    wb = Workbook()
    apply_default_font(wb)
    ws = wb.active
    ws.title = (sheet_name or src.stem)[:31]
    fmts = {}
    for col in df.columns:
        if str(df[col].dtype).startswith("datetime"):
            fmts[str(col)] = "date"
        elif str(df[col].dtype).startswith("float"):
            fmts[str(col)] = "dec2"
        elif str(df[col].dtype).startswith("int"):
            fmts[str(col)] = "int"
    write_dataframe(ws, df, number_formats=fmts)
    ws.auto_filter.ref = ws.dimensions
    wb.save(dst)
    return {"rows": len(df), "columns": len(df.columns), "sep": sep, "encoding": enc}


def xlsx_to_csv(src: Path, dst: Path, sheet: str | None, sep: str, all_sheets: bool):
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    sheets = wb.sheetnames if all_sheets else [sheet or wb.sheetnames[0]]
    outputs = []
    for name in sheets:
        ws = wb[name]
        out = (dst / f"{name}.csv") if all_sheets else dst
        out.parent.mkdir(parents=True, exist_ok=True)
        with open(out, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=sep)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
        outputs.append(str(out))
    return outputs


def lo_convert(src: Path, dst: Path, timeout: int):
    target = LO_TARGETS.get(dst.suffix.lower())
    if not target:
        raise SystemExit(f"Destino {dst.suffix} no soportado")
    work = Path(tempfile.mkdtemp(prefix="conv_"))
    try:
        produced = _soffice.convert(src.resolve(), work, target, timeout=timeout)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), dst)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--to", help="Forzar formato destino (csv) cuando dst es carpeta")
    ap.add_argument("--sheet")
    ap.add_argument("--sep")
    ap.add_argument("--decimal")
    ap.add_argument("--encoding")
    ap.add_argument("--timeout", type=int, default=120)
    a = ap.parse_args()
    src, dst = Path(a.src), Path(a.dst)
    s_ext, d_ext = src.suffix.lower(), (("." + a.to) if a.to else dst.suffix.lower())

    if s_ext in (".csv", ".tsv", ".txt") and d_ext == ".xlsx":
        info = csv_to_xlsx(src, dst, a.sep or ("\t" if s_ext == ".tsv" else None), a.decimal,
                           a.encoding, a.sheet)
        print(f"OK → {dst}  {info}")
    elif d_ext == ".csv":
        if s_ext not in (".xlsx", ".xlsm"):
            tmp = Path(tempfile.mkdtemp()) / (src.stem + ".xlsx")
            lo_convert(src, tmp, a.timeout)
            src = tmp
        outs = xlsx_to_csv(src, dst, a.sheet, a.sep or ",", all_sheets=bool(a.to) or dst.suffix == "")
        print("OK →", outs)
    else:
        lo_convert(src, dst, a.timeout)
        print(f"OK → {dst}")


if __name__ == "__main__":
    main()
