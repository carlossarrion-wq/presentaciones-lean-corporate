#!/usr/bin/env python3
"""Auditoría estática de fórmulas: encuentra fórmulas que EVALÚAN sin error pero
probablemente están MAL. Complementa a recalc.py (que solo detecta errores #...).

Detecta:
  inconsistent   Fórmula distinta a sus vecinas de columna/fila (p.ej. una celda editada a
                 mano dentro de una fila de proyecciones). Comparación en notación R1C1.
  hardcoded      Números "mágicos" dentro de fórmulas (=B5*1.21) en lugar de una celda de supuesto.
  empty_ref      Referencias a celdas vacías (posible rango desplazado).
  unguarded_div  Divisiones cuyo denominador es una referencia, sin IFERROR/IF que la proteja.
  self_range     Rangos que incluyen la propia celda (referencia circular).

Uso:
    python audit_formulas.py libro.xlsx [--sheet Hoja] [--allow 0,1,12,100] [--json]
Los números permitidos por defecto (0, 1, 2, 12, 100, 1000) no se marcan como mágicos.
"""
from __future__ import annotations

import argparse
import json
import re
import warnings
from collections import Counter, defaultdict

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
AGG = re.compile(r"^=\s*(SUM|AVERAGE|SUBTOTAL|AGGREGATE|COUNT|COUNTA|MIN|MAX|MEDIAN|SUMPRODUCT|"
                 r"_xlfn\.AGGREGATE)\(", re.I)


def to_r1c1(formula: str, row: int, col: int) -> str:
    out = []
    for t in Tokenizer(formula).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            def repl(m):
                dc, c, dr, r = m.groups()
                ci, ri = column_index_from_string(c), int(r)
                rs = f"R{ri}" if dr else f"R[{ri - row}]"
                cs = f"C{ci}" if dc else f"C[{ci - col}]"
                return rs + cs
            out.append(REF.sub(repl, t.value))
        else:
            out.append(t.value)
    return "".join(out)


def audit(path, only_sheet=None, allowed=(0, 1, 2, 12, 100, 1000)):
    wb = openpyxl.load_workbook(path)
    findings = []
    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        formulas = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas[(c.row, c.column)] = c.value
        r1 = {k: to_r1c1(v, *k) for k, v in formulas.items()}

        # 1) inconsistencias: compara con vecinos inmediatos en columna y en fila
        for (r, cidx), pat in r1.items():
            for (dr, dc), axis in (((1, 0), "columna"), ((0, 1), "fila")):
                a, b = (r - dr, cidx - dc), (r + dr, cidx + dc)
                a2, b2 = (r - 2 * dr, cidx - 2 * dc), (r + 2 * dr, cidx + 2 * dc)
                middle = a in r1 and b in r1 and r1[a] == r1[b] and r1[a] != pat
                # extremos de un bloque: los dos anteriores (o siguientes) coinciden entre sí
                tail = b not in r1 and a in r1 and a2 in r1 and r1[a] == r1[a2] and r1[a] != pat
                head = a not in r1 and b in r1 and b2 in r1 and r1[b] == r1[b2] and r1[b] != pat
                if (tail or head) and AGG.match(formulas[(r, cidx)]):
                    tail = head = False  # fila/columna de totales al final de un bloque: normal
                if tail or head:
                    a = a if tail else b
                if middle or tail or head:
                    findings.append({"type": "inconsistent", "cell": f"{ws.title}!{get_column_letter(cidx)}{r}",
                                     "detail": f"Distinta de sus vecinas en la {axis}: {formulas[(r, cidx)]} "
                                               f"vs patrón {formulas[a]} (en {get_column_letter(a[1])}{a[0]})"})
                    break

        for (r, cidx), f in formulas.items():
            ref = f"{ws.title}!{get_column_letter(cidx)}{r}"
            toks = Tokenizer(f).items
            # 2) números mágicos
            magic = [t.value for t in toks if t.type == Token.OPERAND and t.subtype == Token.NUMBER
                     and float(t.value) not in allowed]
            if magic:
                findings.append({"type": "hardcoded", "cell": ref,
                                 "detail": f"{f}  → constantes {magic}; muévelas a una celda de supuesto"})
            # 3) referencias a celdas vacías (solo celdas sueltas de la misma hoja)
            for t in toks:
                if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" not in t.value \
                        and ":" not in t.value and REF.fullmatch(t.value):
                    coord = t.value.replace("$", "")
                    if ws[coord].value is None:
                        findings.append({"type": "empty_ref", "cell": ref,
                                         "detail": f"{f} → {coord} está vacía"})
                # 5) rango que contiene la propia celda
                if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" not in t.value \
                        and ":" in t.value and REF.search(t.value):
                    try:
                        c1, r1_, c2, r2 = range_boundaries(t.value.replace("$", ""))
                        if c1 <= cidx <= c2 and r1_ <= r <= r2:
                            findings.append({"type": "self_range", "cell": ref,
                                             "detail": f"{f} incluye su propia celda (circular)"})
                    except Exception:
                        pass
            # 4) divisiones sin protección
            up = f.upper()
            if "/" in f and "IFERROR" not in up and not re.search(r"\bIF\(", up):
                for i, t in enumerate(toks):
                    if t.type == Token.OP_IN and t.value == "/" and i + 1 < len(toks):
                        nxt = toks[i + 1]
                        if nxt.type == Token.OPERAND and nxt.subtype == Token.RANGE:
                            findings.append({"type": "unguarded_div", "cell": ref,
                                             "detail": f"{f} → divisor {nxt.value} sin IFERROR/IF; "
                                                       "dará #DIV/0! si es 0 o vacío"})
                            break
    return findings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--allow", default="0,1,2,12,100,1000")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--max", type=int, default=60)
    a = ap.parse_args()
    allowed = tuple(float(x) for x in a.allow.split(",") if x.strip())
    res = audit(a.file, a.sheet, allowed)
    if a.json:
        print(json.dumps({"total": len(res), "by_type": Counter(f["type"] for f in res), "findings": res},
                         ensure_ascii=False, indent=2))
        return
    if not res:
        print("Sin hallazgos.")
        return
    by = defaultdict(list)
    for f in res:
        by[f["type"]].append(f)
    for t, items in by.items():
        print(f"\n## {t} ({len(items)})")
        for f in items[: a.max]:
            print(f"  {f['cell']}: {f['detail']}")
        if len(items) > a.max:
            print(f"  … y {len(items) - a.max} más")


if __name__ == "__main__":
    main()
