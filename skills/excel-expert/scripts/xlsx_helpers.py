"""Biblioteca de ayudantes para openpyxl: estilos profesionales, tablas, gráficos,
validaciones, formato condicional, volcado de DataFrames e inserción de filas que SÍ
actualiza las referencias de las fórmulas.

Importación desde un script del agente:
    import sys; sys.path.insert(0, "<ruta-del-skill>/scripts")
    from xlsx_helpers import *

Todas las funciones operan sobre objetos openpyxl y no guardan el libro: el llamante
decide cuándo hacer wb.save(...).
"""
from __future__ import annotations

import re
from typing import Iterable, Sequence

from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import (CellIsRule, ColorScaleRule, DataBarRule,
                                      FormulaRule)
from openpyxl.formula.tokenizer import Tokenizer, Token
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------- estilo base
FONT_NAME = "Arial"
PALETTE = {
    "header_fill": "1F3864",   # azul marino
    "header_font": "FFFFFF",
    "band": "F2F2F2",
    "border": "BFBFBF",
    "input_font": "0000FF",    # azul: datos introducidos a mano
    "link_font": "008000",     # verde: vínculos a otra hoja
    "ext_font": "FF0000",      # rojo: vínculos a otro fichero
    "input_fill": "FFFF00",    # amarillo: supuestos clave / celdas a rellenar
    "good": "C6EFCE", "good_font": "006100",
    "bad": "FFC7CE", "bad_font": "9C0006",
    "neutral": "FFEB9C", "neutral_font": "9C5700",
}

NUMBER_FORMATS = {
    "int": "#,##0",
    "dec2": "#,##0.00",
    "eur": '#,##0 "€";-#,##0 "€";"-"',
    "eur2": '#,##0.00 "€";-#,##0.00 "€";"-"',
    "usd": '$#,##0;($#,##0);"-"',
    "usd2": '$#,##0.00;($#,##0.00);"-"',
    "pct": "0.0%",
    "pct0": "0%",
    "multiple": '0.0"x"',
    "date": "dd/mm/yyyy",
    "date_iso": "yyyy-mm-dd",
    "datetime": "dd/mm/yyyy hh:mm",
    "month": "mmm-yy",
    "thousands": '#,##0,"k"',
    "millions": '#,##0.0,,"M"',
    "text": "@",
}

THIN = Side(style="thin", color=PALETTE["border"])
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(bold=False, color="000000", size=10, italic=False, name=FONT_NAME):
    return Font(name=name, bold=bold, color=color, size=size, italic=italic)


def fill(hex_color: str):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def apply_default_font(wb, name=FONT_NAME, size=10):
    """Fija la fuente por defecto del libro (afecta a celdas sin fuente explícita)."""
    try:
        wb._named_styles["Normal"].font = Font(name=name, size=size)
    except Exception:
        pass


def iter_range(ws, ref: str):
    for row in ws[ref]:
        for c in (row if isinstance(row, tuple) else (row,)):
            yield c


# --------------------------------------------------------------------------- escritura
def write_rows(ws, rows: Iterable[Sequence], start_row=1, start_col=1):
    """Escribe una lista de filas. Devuelve la última fila escrita."""
    r = start_row - 1
    for r, row in enumerate(rows, start=start_row):
        for j, v in enumerate(row):
            ws.cell(row=r, column=start_col + j, value=v)
    return r


def write_dataframe(ws, df, start_row=1, start_col=1, index=False, header=True,
                    table_name: str | None = None, number_formats: dict | None = None,
                    style_header=True, autofit=True, freeze=True):
    """Vuelca un pandas.DataFrame con cabecera estilizada, formatos por columna y,
    opcionalmente, como Tabla de Excel. Devuelve el rango ocupado (p.ej. 'A1:F120')."""
    import pandas as pd  # import diferido
    data = df.reset_index() if index else df
    cols = [str(c) for c in data.columns]
    r = start_row
    if header:
        for j, name in enumerate(cols):
            ws.cell(row=r, column=start_col + j, value=name)
        r += 1
    for rec in data.itertuples(index=False, name=None):
        for j, v in enumerate(rec):
            if v is not None and isinstance(v, float) and v != v:  # NaN → vacío
                v = None
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime() if not pd.isna(v) else None
            elif v is pd.NaT:
                v = None
            ws.cell(row=r, column=start_col + j, value=v)
        r += 1
    last_row = max(r - 1, start_row)
    last_col = start_col + len(cols) - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(last_col)}{last_row}"
    for j, name in enumerate(cols):
        fmt = (number_formats or {}).get(name)
        if fmt:
            fmt = NUMBER_FORMATS.get(fmt, fmt)
            for i in range(start_row + (1 if header else 0), last_row + 1):
                ws.cell(row=i, column=start_col + j).number_format = fmt
    if header and style_header:
        # explícito incluso con Tabla: LibreOffice/Google Sheets no renderizan estilos de tabla
        style_header_row(ws, start_row, start_col, last_col)
    if table_name and header and last_row > start_row:
        add_table(ws, ref, table_name)
    if freeze and header and start_col == 1:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    if autofit:
        autofit_columns(ws)
    return ref


# --------------------------------------------------------------------------- estilos
def style_header_row(ws, row=1, start_col=1, end_col=None, fill_color=None,
                     font_color=None, wrap=True, height=None):
    end_col = end_col or ws.max_column
    for col in range(start_col, end_col + 1):
        c = ws.cell(row=row, column=col)
        c.font = font(bold=True, color=font_color or PALETTE["header_font"])
        c.fill = fill(fill_color or PALETTE["header_fill"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        c.border = BORDER_ALL
    if height:
        ws.row_dimensions[row].height = height


def style_range(ws, ref, *, number_format=None, bold=None, color=None, fill_color=None,
                border=False, align=None, wrap=None, font_name=FONT_NAME, size=None):
    for c in iter_range(ws, ref):
        if number_format:
            c.number_format = NUMBER_FORMATS.get(number_format, number_format)
        if any(x is not None for x in (bold, color, size)) or font_name:
            f = c.font
            c.font = Font(name=font_name or f.name, size=size or f.size or 10,
                          bold=f.bold if bold is None else bold,
                          italic=f.italic, color=color or (f.color if f.color else None))
        if fill_color:
            c.fill = fill(fill_color)
        if border:
            c.border = BORDER_ALL
        if align or wrap is not None:
            c.alignment = Alignment(horizontal=align or c.alignment.horizontal,
                                    vertical="center", wrap_text=wrap)


def band_rows(ws, ref, color=None):
    """Bandas alternas (cebra) en un rango."""
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    for r in range(min_row, max_row + 1):
        if (r - min_row) % 2 == 1:
            for col in range(min_col, max_col + 1):
                ws.cell(r, col).fill = fill(color or PALETTE["band"])


def mark_input(cell, key_assumption=False):
    """Celda de entrada: fuente azul (y relleno amarillo si es supuesto clave)."""
    cell.font = font(color=PALETTE["input_font"], bold=cell.font.bold)
    if key_assumption:
        cell.fill = fill(PALETTE["input_fill"])


def autofit_columns(ws, min_width=8, max_width=60, padding=2):
    """Ajusta el ancho de columna al contenido (aproximado; Excel no guarda 'autofit').
    Las fórmulas se miden como 12 caracteres al no conocer su resultado."""
    widths: dict[int, float] = {}
    merged = set()
    for mr in ws.merged_cells.ranges:
        for row in ws.iter_rows(min_row=mr.min_row, max_row=mr.max_row,
                                min_col=mr.min_col, max_col=mr.max_col):
            for c in row:
                merged.add(c.coordinate)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.coordinate in merged:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                length = 12
            elif isinstance(v, float):
                length = len(f"{v:,.2f}")
            elif hasattr(v, "strftime"):
                length = 10
            else:
                length = max(len(line) for line in str(v).split("\n"))
            if c.font is not None and c.font.bold:
                length *= 1.1
            widths[c.column] = max(widths.get(c.column, 0), length)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + padding))


def set_column_widths(ws, widths: dict):
    """widths = {'A': 30, 'B': 12, ...}"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def freeze_header(ws, cell="A2"):
    ws.freeze_panes = cell


def add_table(ws, ref, name, style="TableStyleMedium2", striped=True):
    """Convierte un rango (con cabecera en la 1ª fila) en Tabla de Excel con filtros.
    El nombre debe ser único en el libro, sin espacios y no parecer una referencia (A1)."""
    name = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not re.match(r"[A-Za-z_]", name):
        name = "T_" + name
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False,
                                      showRowStripes=striped, showColumnStripes=False)
    ws.add_table(t)
    return t


def add_comment(cell, text, author="Agente IA", width=300, height=100):
    c = Comment(text, author)
    c.width, c.height = width, height
    cell.comment = c


def add_hyperlink_to_sheet(cell, sheet_name, target_cell="A1", text=None):
    cell.hyperlink = f"#'{sheet_name}'!{target_cell}"
    cell.value = text or sheet_name
    cell.font = Font(name=FONT_NAME, color="0563C1", underline="single")


def set_print_setup(ws, landscape=True, fit_width=True, repeat_header_rows="1:1", paper_a4=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    if paper_a4:
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
    if fit_width:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    if repeat_header_rows:
        ws.print_title_rows = repeat_header_rows


# --------------------------------------------------------------------------- validaciones
def add_list_validation(ws, sqref, options: Sequence[str] | None = None, source_range: str | None = None,
                        allow_blank=True, error_msg="Valor no permitido"):
    """Desplegable. options=['Sí','No'] (máx. 255 caracteres en total) o
    source_range="'Listas'!$A$2:$A$20" para listas largas."""
    if options:
        formula = '"' + ",".join(str(o).replace('"', '""') for o in options) + '"'
        if len(formula) > 255:
            raise ValueError("Lista demasiado larga para validación inline; usa source_range")
    elif source_range:
        formula = source_range if source_range.startswith("=") else "=" + source_range
    else:
        raise ValueError("Indica options o source_range")
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank,
                        showErrorMessage=True, errorTitle="Entrada no válida", error=error_msg)
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def add_number_validation(ws, sqref, minimum=None, maximum=None, whole=False):
    op = "between" if minimum is not None and maximum is not None else (
        "greaterThanOrEqual" if minimum is not None else "lessThanOrEqual")
    dv = DataValidation(type="whole" if whole else "decimal", operator=op,
                        formula1=str(minimum if minimum is not None else maximum),
                        formula2=str(maximum) if op == "between" else None,
                        showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


# --------------------------------------------------------------------------- formato condicional
def add_color_scale(ws, ref, start="F8696B", mid="FFEB84", end="63BE7B"):
    ws.conditional_formatting.add(ref, ColorScaleRule(start_type="min", start_color=start,
                                                      mid_type="percentile", mid_value=50, mid_color=mid,
                                                      end_type="max", end_color=end))


def add_data_bars(ws, ref, color="638EC6"):
    ws.conditional_formatting.add(ref, DataBarRule(start_type="min", end_type="max", color=color))


def highlight_cells(ws, ref, operator, values: Sequence, kind="bad"):
    """operator: 'lessThan','greaterThan','between','equal'... kind: good|bad|neutral"""
    ws.conditional_formatting.add(ref, CellIsRule(
        operator=operator, formula=[str(v) for v in values],
        fill=fill(PALETTE[kind]), font=Font(color=PALETTE[f"{kind}_font"])))


def highlight_formula(ws, ref, formula, kind="neutral"):
    """Regla por fórmula relativa a la celda superior izquierda de ref, p.ej. '$E2>1000'."""
    ws.conditional_formatting.add(ref, FormulaRule(
        formula=[formula.lstrip("=")], fill=fill(PALETTE[kind]),
        font=Font(color=PALETTE[f"{kind}_font"])))


# --------------------------------------------------------------------------- gráficos
def _chart_common(chart, title, x_title, y_title, width, height, style):
    chart.title = title
    if x_title is not None and hasattr(chart, "x_axis"):
        chart.x_axis.title = x_title
    if y_title is not None and hasattr(chart, "y_axis"):
        chart.y_axis.title = y_title
    chart.width, chart.height = width, height
    if style:
        chart.style = style
    # openpyxl ≥3.1 oculta ejes por defecto en algunos visores: forzarlos visibles
    for ax in ("x_axis", "y_axis"):
        if hasattr(chart, ax):
            getattr(chart, ax).delete = False
    return chart


def add_bar_chart(ws, data_ref: str, cats_ref: str, anchor: str, title="", x_title=None,
                  y_title=None, horizontal=False, stacked=False, width=18, height=9,
                  titles_from_data=True, data_sheet=None, style=10):
    """data_ref incluye la fila de cabecera (nombres de series) si titles_from_data=True.
    Ej.: add_bar_chart(ws, 'B1:C13', 'A2:A13', 'E2', 'Ventas por mes')"""
    src = data_sheet or ws
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    c1, r1, c2, r2 = range_boundaries(data_ref)
    ch.add_data(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2),
                titles_from_data=titles_from_data)
    c1, r1, c2, r2 = range_boundaries(cats_ref)
    ch.set_categories(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
    ws.add_chart(_chart_common(ch, title, x_title, y_title, width, height, style), anchor)
    return ch


def add_line_chart(ws, data_ref, cats_ref, anchor, title="", x_title=None, y_title=None,
                   width=18, height=9, titles_from_data=True, data_sheet=None, style=12, markers=True):
    src = data_sheet or ws
    ch = LineChart()
    c1, r1, c2, r2 = range_boundaries(data_ref)
    ch.add_data(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2),
                titles_from_data=titles_from_data)
    c1, r1, c2, r2 = range_boundaries(cats_ref)
    ch.set_categories(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
    if markers:
        for s in ch.series:
            s.marker.symbol = "circle"
            s.smooth = False
    ws.add_chart(_chart_common(ch, title, x_title, y_title, width, height, style), anchor)
    return ch


def add_pie_chart(ws, data_ref, cats_ref, anchor, title="", width=12, height=9,
                  titles_from_data=True, data_sheet=None, show_percent=True):
    from openpyxl.chart.label import DataLabelList
    src = data_sheet or ws
    ch = PieChart()
    c1, r1, c2, r2 = range_boundaries(data_ref)
    ch.add_data(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2),
                titles_from_data=titles_from_data)
    c1, r1, c2, r2 = range_boundaries(cats_ref)
    ch.set_categories(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
    if show_percent:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
    ch.title, ch.width, ch.height = title, width, height
    ws.add_chart(ch, anchor)
    return ch


def add_scatter_chart(ws, x_ref, y_refs: Sequence[str], anchor, title="", x_title=None,
                      y_title=None, series_titles: Sequence[str] | None = None,
                      width=18, height=9, data_sheet=None):
    src = data_sheet or ws
    ch = ScatterChart()
    c1, r1, c2, r2 = range_boundaries(x_ref)
    xs = Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
    for i, yref in enumerate(y_refs):
        c1, r1, c2, r2 = range_boundaries(yref)
        s = Series(Reference(src, min_col=c1, min_row=r1, max_col=c2, max_row=r2), xs,
                   title=(series_titles[i] if series_titles else None))
        s.marker.symbol = "circle"
        s.graphicalProperties.line.noFill = True
        ch.series.append(s)
    ws.add_chart(_chart_common(ch, title, x_title, y_title, width, height, None), anchor)
    return ch


# --------------------------------------------------------------------------- inserción segura
_CELL_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_REF_RE = re.compile(r"\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}|\$?\d+:\$?\d+")


def _shift_ref_text(ref: str, row_at: int, n_rows: int, col_at: int, n_cols: int) -> str:
    def repl(m):
        dc, col, dr, row = m.groups()
        row_i, col_i = int(row), column_index_from_string(col)
        if n_rows and row_i >= row_at:
            row_i += n_rows
        if n_cols and col_i >= col_at:
            col_i += n_cols
        return f"{dc}{get_column_letter(col_i)}{dr}{row_i}"
    # rangos de columna completa (A:C) o fila completa (3:5)
    if re.fullmatch(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", ref) and n_cols:
        parts = []
        for p in ref.split(":"):
            d = "$" if p.startswith("$") else ""
            ci = column_index_from_string(p.lstrip("$"))
            parts.append(d + get_column_letter(ci + n_cols if ci >= col_at else ci))
        return ":".join(parts)
    if re.fullmatch(r"\$?\d+:\$?\d+", ref) and n_rows:
        parts = []
        for p in ref.split(":"):
            d = "$" if p.startswith("$") else ""
            ri = int(p.lstrip("$"))
            parts.append(d + str(ri + n_rows if ri >= row_at else ri))
        return ":".join(parts)
    return _CELL_RE.sub(repl, ref)


def _shift_formula(formula: str, target_sheet: str, formula_sheet: str,
                   row_at: int, n_rows: int, col_at: int, n_cols: int) -> str:
    tok = Tokenizer(formula)
    out = []
    for t in tok.items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            v = t.value
            if "!" in v:
                sheet_part, ref = v.rsplit("!", 1)
                sname = sheet_part.strip("'").replace("''", "'")
                if sname == target_sheet and not sheet_part.startswith("[") and _REF_RE.fullmatch(ref):
                    v = sheet_part + "!" + _shift_ref_text(ref, row_at, n_rows, col_at, n_cols)
            elif formula_sheet == target_sheet and _REF_RE.fullmatch(v):
                v = _shift_ref_text(v, row_at, n_rows, col_at, n_cols)
            out.append(v)
        else:
            out.append(t.value)
    return "=" + "".join(out)


def _shift_structures(wb, ws, row_at, n_rows, col_at, n_cols):
    title = ws.title
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = _shift_formula(c.value, title, sheet.title, row_at, n_rows, col_at, n_cols)
    # nombres definidos del libro
    try:
        for name, dn in list(wb.defined_names.items()):
            txt = dn.attr_text
            if txt and title in txt:
                dn.attr_text = _shift_formula("=" + txt, title, "", row_at, n_rows, col_at, n_cols)[1:]
    except Exception:
        pass
    # celdas combinadas
    new_ranges = []
    for mr in list(ws.merged_cells.ranges):
        new_ranges.append(_shift_ref_text(mr.coord, row_at, n_rows, col_at, n_cols))
        ws.unmerge_cells(mr.coord)
    # tablas
    for t in ws.tables.values():
        t.ref = _shift_ref_text(t.ref, row_at, n_rows, col_at, n_cols)
        if t.autoFilter is not None and t.autoFilter.ref:
            t.autoFilter.ref = t.ref
    # validaciones y formato condicional
    for dv in ws.data_validations.dataValidation:
        dv.sqref = " ".join(_shift_ref_text(str(r), row_at, n_rows, col_at, n_cols)
                            for r in str(dv.sqref).split())
    from openpyxl.formatting.formatting import ConditionalFormattingList
    old_cf = ws.conditional_formatting
    new_cf = ConditionalFormattingList()
    for cf in old_cf:
        new_ref = " ".join(_shift_ref_text(str(r), row_at, n_rows, col_at, n_cols)
                           for r in str(cf.sqref).split())
        for rule in cf.rules:
            new_cf.add(new_ref, rule)
    ws.conditional_formatting = new_cf
    return new_ranges


def insert_rows_safe(wb, ws, idx: int, amount: int = 1, copy_style_from: int | None = None):
    """Inserta `amount` filas antes de la fila `idx` y ACTUALIZA: fórmulas de todo el libro
    que apunten a esta hoja, nombres definidos, celdas combinadas, tablas, validaciones y
    formato condicional. (ws.insert_rows de openpyxl NO actualiza nada de eso.)

    copy_style_from: fila (numeración previa a la inserción) cuyo estilo y fórmulas
    (trasladadas) se copian a las filas nuevas. Útil para ampliar una tabla.
    Limitación: gráficos e impresión no se ajustan; revisa sus rangos si dependen de la zona.
    """
    from copy import copy
    from openpyxl.formula.translate import Translator
    new_merged = _shift_structures(wb, ws, idx, amount, 10**6, 0)
    ws.insert_rows(idx, amount)
    for rng in new_merged:
        ws.merge_cells(rng)
    if copy_style_from is not None:
        src_row = copy_style_from + (amount if copy_style_from >= idx else 0)
        for r in range(idx, idx + amount):
            for col in range(1, ws.max_column + 1):
                s = ws.cell(src_row, col)
                d = ws.cell(r, col)
                if s.has_style:
                    d._style = copy(s._style)
                if isinstance(s.value, str) and s.value.startswith("="):
                    d.value = Translator(s.value, origin=s.coordinate).translate_formula(d.coordinate)
            ws.row_dimensions[r].height = ws.row_dimensions[src_row].height


def insert_cols_safe(wb, ws, idx: int, amount: int = 1):
    """Igual que insert_rows_safe pero para columnas (antes de la columna nº idx)."""
    new_merged = _shift_structures(wb, ws, 10**7, 0, idx, amount)
    ws.insert_cols(idx, amount)
    for rng in new_merged:
        ws.merge_cells(rng)


__all__ = [n for n in dir() if not n.startswith("_") and n not in
           ("annotations", "re", "Iterable", "Sequence")]
