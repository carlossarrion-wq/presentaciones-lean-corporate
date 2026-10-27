#!/usr/bin/env python3
"""Recalcula todas las fórmulas de un libro con LibreOffice y detecta errores.

openpyxl escribe las fórmulas como texto SIN valor calculado. Hasta recalcular, cualquier
lector de valores (pandas, openpyxl data_only=True, previsualizadores) ve `None`.

Uso:
    python recalc.py libro.xlsx                 # recalcula y sobrescribe (in situ)
    python recalc.py libro.xlsx -o salida.xlsx  # recalcula a otro fichero
    python recalc.py libro.xlsx --check-only    # solo valida; NO modifica nada
    python recalc.py libro.xlsx --timeout 120

Salida: JSON por stdout.
    status: "success" | "errors_found"
    total_formulas, total_errors, error_summary{tipo: {count, locations[]}},
    warnings[] (problemas detectados antes del recálculo), empty_results (fórmulas que
    devuelven vacío).
Si aparece la clave "error" en lugar de "status", NO se recalculó nada (exit code 1).
"errors_found" sale con código 0 a propósito: revisa siempre el JSON, no solo el exit code.
"""
from __future__ import annotations

import argparse
import json
import warnings
import re
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _soffice  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print(json.dumps({"error": "Falta openpyxl: pip install openpyxl"}))
    sys.exit(1)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ERROR_VALUES = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
                "#SPILL!", "#CALC!", "#GETTING_DATA", "Err:")
MAX_LOCATIONS = 100

# Funciones que LibreOffice (24.x) no evalúa bien desde un fichero escrito por openpyxl
# (funciones matriciales dinámicas sin metadatos de desbordamiento).
UNSAFE_FUNCS = ["XLOOKUP", "XMATCH", "SORT", "SORTBY", "FILTER", "UNIQUE", "SEQUENCE",
                "RANDARRAY", "LET", "LAMBDA", "MAP", "REDUCE", "SCAN", "BYROW", "BYCOL",
                "MAKEARRAY", "VSTACK", "HSTACK", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS",
                "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "TOCOL", "TOROW", "WRAPROWS",
                "WRAPCOLS", "EXPAND", "IMAGE"]
# Funciones posteriores a Excel 2007 que necesitan prefijo _xlfn. en el XML.
NEEDS_XLFN = ["TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS", "IFNA",
              "XOR", "DAYS", "ISOWEEKNUM", "STDEV.S", "STDEV.P", "VAR.S", "VAR.P",
              "PERCENTILE.INC", "PERCENTILE.EXC", "QUARTILE.INC", "QUARTILE.EXC",
              "MODE.SNGL", "RANK.EQ", "RANK.AVG", "CEILING.MATH", "FLOOR.MATH",
              "NORM.DIST", "NORM.INV", "NORM.S.DIST", "NORM.S.INV", "T.TEST", "T.DIST",
              "CHISQ.TEST", "CONFIDENCE.NORM", "COVARIANCE.S", "COVARIANCE.P",
              "FORECAST.LINEAR", "AGGREGATE", "NETWORKDAYS.INTL", "WORKDAY.INTL"]


def _strip_strings(formula: str) -> str:
    return re.sub(r'"[^"]*"', '""', formula)


def preflight(path: Path) -> tuple[list[str], int, bool, list[tuple[str, str]]]:
    """Analiza las fórmulas sin recalcular. Devuelve (avisos, nº fórmulas, hay_links, celdas)."""
    warnings: list[str] = []
    keep_vba = path.suffix.lower() in (".xlsm", ".xltm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    sheet_names = set(wb.sheetnames)
    formula_cells: list[tuple[str, str]] = []
    unsafe_hits: dict[str, list[str]] = {}
    xlfn_hits: dict[str, list[str]] = {}
    bad_sheet_refs: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    ref = f"{ws.title}!{c.coordinate}"
                    formula_cells.append((ws.title, c.coordinate))
                    f = _strip_strings(v).upper()
                    for fn in UNSAFE_FUNCS:
                        if re.search(rf"(?<![A-Z0-9_.]){re.escape(fn)}\(", f):
                            unsafe_hits.setdefault(fn, []).append(ref)
                    for fn in NEEDS_XLFN:
                        if re.search(rf"(?<![A-Z0-9_.]){re.escape(fn)}\(", f):
                            xlfn_hits.setdefault(fn, []).append(ref)
                    # Referencias a hojas: 'Hoja X'!A1 o Hoja!A1
                    for m in re.finditer(r"'((?:[^']|'')+)'!|(?<![A-Z0-9_\]'])([A-Za-z_][\w.]*)!",
                                         _strip_strings(v)):
                        name = (m.group(1) or m.group(2) or "").replace("''", "'")
                        if name and name not in sheet_names and not name.startswith("["):
                            bad_sheet_refs.append(f"{ref} -> '{name}'")
                elif type(v).__name__ == "ArrayFormula":
                    formula_cells.append((ws.title, c.coordinate))
    for fn, refs in unsafe_hits.items():
        warnings.append(f"{fn} no es fiable fuera de Excel 365 (matriz dinámica); usa "
                        f"INDEX/MATCH o calcula en Python. Celdas: {refs[:10]}")
    for fn, refs in xlfn_hits.items():
        warnings.append(f"{fn} debe escribirse como _xlfn.{fn}(...) o dará #NAME?. "
                        f"Celdas: {refs[:10]}")
    if bad_sheet_refs:
        warnings.append("Referencias a hojas inexistentes o nombres con espacios sin "
                        f"comillas simples: {bad_sheet_refs[:10]}")
    has_links = bool(getattr(wb, "_external_links", None))
    if has_links:
        warnings.append("El libro contiene vínculos a ficheros externos ([1]Hoja!A1). "
                        "Al recalcular sin esos ficheros, los valores se pierden (#REF!/#NAME?).")
    return warnings, len(formula_cells), has_links, formula_cells


def scan_errors(recalculated: Path, formula_cells: list[tuple[str, str]]) -> dict:
    wb = openpyxl.load_workbook(recalculated, data_only=True)
    summary: dict[str, dict] = {}
    empty = 0
    total_errors = 0
    for sheet, coord in formula_cells:
        if sheet not in wb.sheetnames:
            continue
        val = wb[sheet][coord].value
        if val is None:
            empty += 1
            continue
        if isinstance(val, str) and val.startswith(ERROR_VALUES):
            key = val if not val.startswith("Err:") else "Err:LibreOffice"
            entry = summary.setdefault(key, {"count": 0, "locations": []})
            entry["count"] += 1
            total_errors += 1
            if len(entry["locations"]) < MAX_LOCATIONS:
                entry["locations"].append(f"{sheet}!{coord}")
    for entry in summary.values():
        withheld = entry["count"] - len(entry["locations"])
        if withheld:
            entry["locations_truncated"] = withheld
    return {"total_errors": total_errors, "error_summary": summary, "empty_results": empty}


def target_spec(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".xlsm":
        return 'xlsm:Calc MS Excel 2007 VBA XML'
    if ext == ".ods":
        return "ods"
    return 'xlsx:Calc Office Open XML'


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("timeout_pos", nargs="?", type=int, help=argparse.SUPPRESS)
    ap.add_argument("--timeout", type=int, default=60)
    ap.add_argument("-o", "--output", help="Escribir el resultado aquí en vez de sobrescribir")
    ap.add_argument("--check-only", action="store_true", help="Validar sin modificar el fichero")
    ap.add_argument("--force", action="store_true",
                    help="Recalcular aunque haya vínculos externos (se perderán)")
    args = ap.parse_args()
    timeout = args.timeout_pos or args.timeout

    src = Path(args.file).resolve()
    if not src.exists():
        print(json.dumps({"error": f"No existe: {src}"}, ensure_ascii=False))
        return 1
    if src.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm", ".ods"):
        print(json.dumps({"error": "Formato no soportado para recálculo (usa .xlsx/.xlsm/.ods). "
                                   "Convierte antes con convert.py"}, ensure_ascii=False))
        return 1

    try:
        warnings, n_formulas, has_links, cells = preflight(src)
    except Exception as e:  # fichero corrupto, contraseña, etc.
        print(json.dumps({"error": f"No se pudo abrir con openpyxl: {e}"}, ensure_ascii=False))
        return 1

    if has_links and not args.force and not args.check_only:
        print(json.dumps({"error": "Vínculos externos detectados; no se recalcula para no "
                                   "destruir sus valores. Copia esos valores antes o usa --force "
                                   "(o --check-only).", "warnings": warnings}, ensure_ascii=False))
        return 1

    if n_formulas == 0:
        print(json.dumps({"status": "success", "total_formulas": 0, "total_errors": 0,
                          "error_summary": {}, "warnings": warnings,
                          "note": "El libro no contiene fórmulas; nada que recalcular."},
                         ensure_ascii=False, indent=2))
        return 0

    work = Path(tempfile.mkdtemp(prefix="recalc_"))
    try:
        tmp_in = work / "in" / src.name
        tmp_in.parent.mkdir()
        shutil.copy2(src, tmp_in)
        try:
            produced = _soffice.convert(tmp_in, work / "out", target_spec(src), timeout=timeout)
        except RuntimeError as e:
            print(json.dumps({"error": str(e)}, ensure_ascii=False))
            return 1
        result = scan_errors(produced, cells)
        result = {"status": "errors_found" if result["total_errors"] else "success",
                  "total_formulas": n_formulas, **result, "warnings": warnings}
        if not args.check_only:
            dest = Path(args.output).resolve() if args.output else src
            shutil.copy2(produced, dest)
            result["written_to"] = str(dest)
        else:
            result["written_to"] = None
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
