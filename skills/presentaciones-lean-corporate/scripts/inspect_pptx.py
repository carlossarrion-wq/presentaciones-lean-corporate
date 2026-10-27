#!/usr/bin/env python3
"""
inspect_pptx.py — Radiografía de un .pptx/.potx existente.

    python inspect_pptx.py deck.pptx                 # texto por diapositiva (markdown)
    python inspect_pptx.py deck.pptx --shapes        # + posición/tamaño de cada forma
    python inspect_pptx.py plantilla.potx --template # tema (colores, fuentes) y layouts del patrón
    python inspect_pptx.py deck.pptx --json

Úsalo para: (1) entender un deck que hay que editar o resumir, (2) extraer la
identidad visual de una plantilla corporativa y convertirla en un tema propio
(themes/*.json) con --emit-theme.
"""
import argparse
import json
import os
import re
import sys
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation

EMU_IN = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def open_any(path):
    if path.lower().endswith((".potx", ".potm")):
        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False).name
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                zout.writestr(it, data)
        path = tmp
    return Presentation(path)


def theme_info(prs):
    """Lee colores y fuentes del tema del primer patrón."""
    master = prs.slide_masters[0]
    theme_part = None
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            theme_part = rel.target_part
    if theme_part is None:
        return {}
    root = etree.fromstring(theme_part.blob)
    colors = {}
    scheme = root.find(f".//{A}clrScheme")
    if scheme is not None:
        for el in scheme:
            tag = el.tag.replace(A, "")
            c = el[0]
            # sysClr guarda el color real en lastClr (val es 'windowText', 'window'…)
            colors[tag] = c.get("lastClr") if c.tag.endswith("sysClr") else c.get("val")
    fonts = {}
    fs = root.find(f".//{A}fontScheme")
    if fs is not None:
        maj = fs.find(f"{A}majorFont/{A}latin")
        mnr = fs.find(f"{A}minorFont/{A}latin")
        fonts = {"major": maj.get("typeface") if maj is not None else None,
                 "minor": mnr.get("typeface") if mnr is not None else None}
    return {"colors": colors, "fonts": fonts}


def emit_theme(info, name):
    c = info.get("colors", {})
    f = info.get("fonts", {})

    def g(k, d):
        return (c.get(k) or d).upper()
    accents = [g(f"accent{i}", "888888") for i in range(1, 7)]
    return {
        "base": "ejecutivo",
        "description": f"Tema extraído de {name}",
        "bg": g("lt1", "FFFFFF"),
        "bg_dark": g("dk2", g("dk1", "1F1F1F")),
        "text": g("dk1", "1F1F1F"),
        "primary": accents[0],
        "accent": accents[1],
        "accent_2": accents[2],
        "chart": accents,
        "font_heading": f.get("major") or "Calibri",
        "font_body": f.get("minor") or "Calibri",
        "_nota": "Revisa surface, border, text_muted y positive/negative: se heredan de 'base'."
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--shapes", action="store_true")
    ap.add_argument("--template", action="store_true")
    ap.add_argument("--emit-theme", metavar="OUT.json")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    prs = open_any(a.path)
    W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    info = theme_info(prs)
    data = {"file": a.path, "size_in": [round(W, 3), round(H, 3)],
            "theme": info, "layouts": [], "slides": []}
    for m in prs.slide_masters:
        for l in m.slide_layouts:
            data["layouts"].append({"name": l.name, "placeholders": [
                f"{p.placeholder_format.type}".split(".")[-1].split(" ")[0] for p in l.placeholders]})
    for i, s in enumerate(prs.slides, 1):
        sd = {"n": i, "layout": s.slide_layout.name, "shapes": [], "notes": ""}
        for sh in s.shapes:
            item = {"name": sh.name, "type": str(sh.shape_type).split(".")[-1].split(" ")[0]
                    if sh.shape_type else "?"}
            if sh.width:
                item["box_in"] = [round(sh.left / EMU_IN, 2), round(sh.top / EMU_IN, 2),
                                  round(sh.width / EMU_IN, 2), round(sh.height / EMU_IN, 2)]
            if sh.has_text_frame and sh.text_frame.text.strip():
                item["text"] = sh.text_frame.text.strip()
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                item["chart"] = {"type": str(ch.chart_type).split(".")[-1].split(" ")[0],
                                 "categories": [str(c) for c in ch.plots[0].categories],
                                 "series": [{"name": se.name, "values": list(se.values)}
                                            for se in ch.series]}
            if getattr(sh, "has_table", False) and sh.has_table:
                item["table"] = [[c.text for c in r.cells] for r in sh.table.rows]
            sd["shapes"].append(item)
        if s.has_notes_slide:
            sd["notes"] = s.notes_slide.notes_text_frame.text
        data["slides"].append(sd)

    if a.emit_theme:
        with open(a.emit_theme, "w", encoding="utf-8") as fh:
            json.dump(emit_theme(info, os.path.basename(a.path)), fh, indent=2, ensure_ascii=False)
        print(f"Tema escrito en {a.emit_theme}")
    if a.json:
        print(json.dumps(data, ensure_ascii=False, indent=2, default=str))
        return
    print(f"# {os.path.basename(a.path)}  ({W:.2f} × {H:.2f} in, {len(prs.slides)} diapositivas)\n")
    if a.template or not prs.slides:
        print("## Tema")
        for k, v in info.get("colors", {}).items():
            print(f"- {k}: #{v}")
        print(f"- fuentes: títulos={info.get('fonts', {}).get('major')}, "
              f"cuerpo={info.get('fonts', {}).get('minor')}\n")
        print("## Layouts del patrón")
        for l in data["layouts"]:
            print(f"- {l['name']}: {', '.join(l['placeholders']) or '(sin marcadores)'}")
        print()
    for sd in data["slides"]:
        print(f"## Diapositiva {sd['n']}  _(layout: {sd['layout']})_")
        for it in sd["shapes"]:
            if "text" in it:
                pre = f"[{it['name']} @ {it.get('box_in')}] " if a.shapes else ""
                txt = it["text"].replace("\n", " / ")
                print(f"- {pre}{txt}")
            elif "chart" in it:
                print(f"- [gráfico {it['chart']['type']}] categorías={it['chart']['categories']}")
                for se in it["chart"]["series"]:
                    print(f"    · {se['name']}: {se['values']}")
            elif "table" in it:
                print(f"- [tabla {len(it['table'])}×{len(it['table'][0]) if it['table'] else 0}]")
                for r in it["table"][:12]:
                    print("    | " + " | ".join(c.replace("\n", " ") for c in r) + " |")
            elif a.shapes:
                print(f"- [{it['type']} {it['name']} @ {it.get('box_in')}]")
        if sd["notes"].strip():
            print(f"  > Notas: {sd['notes'].strip()[:300]}")
        print()


if __name__ == "__main__":
    main()
