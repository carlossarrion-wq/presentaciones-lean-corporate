#!/usr/bin/env python3
"""
qa_check.py — Revisión automática de calidad de un .pptx (generado o existente).

    python qa_check.py presentacion.pptx
    python qa_check.py presentacion.pptx --json      # salida máquina
    python qa_check.py presentacion.pptx --strict    # código de salida 1 si hay errores

Detecta:
  ERROR   texto que no cabe en su caja · formas fuera del lienzo · texto de relleno
          (Lorem, TODO, XXX, [Imagen…], {{…}}) · cajas de texto vacías
  AVISO   letra < 10 pt en contenido · más de 3 familias tipográficas · diapositivas
          densas (> 90 palabras o > 7 viñetas) · cajas de texto que se solapan ·
          títulos de más de 2 líneas
  INFO    diapositivas sin notas del orador

No sustituye a mirar las imágenes (render_preview.py): complementa esa revisión.
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from textfit import text_height  # noqa: E402

EMU_IN = 914400
PLACEHOLDER_RE = re.compile(r"(lorem ipsum|\bTODO\b|\bTBD\b|XXX|\[imagen|\[image|\{\{|\}\}|"
                            r"click to add|haga clic para|insert text|texto aquí)", re.I)


def iter_shapes(shapes, dx=0, dy=0):
    for sh in shapes:
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # grupo
            yield from iter_shapes(sh.shapes, dx, dy)
        else:
            yield sh


def box(sh):
    x, y = (sh.left or 0) / EMU_IN, (sh.top or 0) / EMU_IN
    w, h = (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN
    rot = (getattr(sh, "rotation", 0) or 0) % 180
    if 45 < rot < 135:  # caja girada ~90º: el área visible es la traspuesta
        cx, cy = x + w / 2, y + h / 2
        x, y, w, h = cx - h / 2, cy - w / 2, h, w
    return x, y, w, h


def para_info(tf):
    """[(texto, size_pt, bold, font, level)] para cada párrafo."""
    out = []
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size, bold, font = None, False, None
        for r in p.runs:
            if r.font.size:
                size = max(size or 0, r.font.size.pt)
            bold = bold or bool(r.font.bold)
            font = font or r.font.name
        out.append((txt, size or 18, bold, font or "Calibri", p.level))
    return out


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    inter = ix * iy
    small = min(aw * ah, bw * bh) or 1
    return inter / small


def check(path, max_words=90, max_bullets=7):
    prs = Presentation(path)
    W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    k = W / 13.333  # escala respecto al lienzo estándar (2.0 en la plantilla Lean)
    issues = []
    fonts = set()

    def add(sev, slide, msg):
        issues.append({"severity": sev, "slide": slide, "message": msg})

    for si, slide in enumerate(prs.slides, 1):
        words, bullets = 0, 0
        text_boxes = []
        for sh in iter_shapes(slide.shapes):
            x, y, w, h = box(sh)
            if w > 0 and h > 0 and (x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02):
                add("ERROR", si, f"'{sh.name}' se sale del lienzo "
                                 f"({x:.2f},{y:.2f},{w:.2f}x{h:.2f} in)")
            if sh.has_text_frame:
                tf = sh.text_frame
                paras = para_info(tf)
                full = " ".join(p[0] for p in paras).strip()
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            fonts.add(r.font.name)
                if not full:
                    if sh.shape_type == 17 or sh.is_placeholder:  # textbox / placeholder vacío
                        add("ERROR", si, f"Caja de texto vacía: '{sh.name}'")
                    continue
                if PLACEHOLDER_RE.search(full):
                    add("ERROR", si, f"Texto de relleno o marcador: '{full[:60]}'")
                words += len(full.split())
                is_bulleted = any(p._p.pPr is not None and p._p.pPr.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
                    for p in tf.paragraphs)
                if is_bulleted:  # solo viñetas de primer nivel (las de 2º nivel usan '–')
                    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                    top = 0
                    for p in tf.paragraphs:
                        bc = p._p.pPr.find(A + "buChar") if p._p.pPr is not None else None
                        if bc is not None and bc.get("char") != "–" and p.text.strip():
                            top += 1
                    bullets = max(bullets, top)
                for txt, size, bold, font, lvl in paras:
                    if txt.strip() and size / k < 10 and h / k > 0.3:
                        add("AVISO", si, f"Letra de {size / k:.0f} pt (equiv. 16:9) en "
                                         f"contenido: '{txt[:40]}'")
                        break
                # desbordamiento (solo si no hay autoajuste de PowerPoint)
                autofit = tf._txBody.bodyPr.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit")
                ml = (tf.margin_left or 0) / EMU_IN + (tf.margin_right or 0) / EMU_IN
                mt = (tf.margin_top or 0) / EMU_IN + (tf.margin_bottom or 0) / EMU_IN
                if w > 0.2 and h > 0.15 and autofit is None and tf.word_wrap is not False:
                    need = 0
                    for txt, size, bold, font, lvl in paras:
                        sp = 0
                        need += text_height([(txt, lvl)], w - ml, font, size, bold)
                    # espacio entre párrafos
                    for p in tf.paragraphs[:-1]:
                        if p.space_after is not None:
                            need += p.space_after.pt / 72
                    if need > (h - mt) * 1.08 + 0.05:
                        add("ERROR", si, f"Posible desbordamiento en '{sh.name}': "
                                         f"'{full[:45]}…' necesita ~{need:.2f} in, caja {h:.2f} in")
                if sh.shape_type == 17 and len(full) > 2:
                    text_boxes.append((sh.name, (x, y, w, h), full))
                if si and y / k < 1.4 and paras and paras[0][1] / k >= 24 and len(paras) == 1:
                    lines = text_height([paras[0][0]], w - ml, paras[0][3], paras[0][1],
                                        paras[0][2]) / (paras[0][1] * 1.2 / 72)
                    if lines > 2.3:
                        add("AVISO", si, f"Título de más de 2 líneas: '{full[:50]}…' (acórtalo)")
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                o = overlap(text_boxes[i][1], text_boxes[j][1])
                if o > 0.15:
                    add("AVISO", si, f"Cajas de texto solapadas: '{text_boxes[i][2][:25]}' / "
                                     f"'{text_boxes[j][2][:25]}' ({o:.0%})")
        if words > max_words:
            add("AVISO", si, f"Diapositiva densa: {words} palabras (límite {max_words})")
        if bullets > max_bullets:
            add("AVISO", si, f"{bullets} viñetas en un mismo bloque (objetivo ≤ 6): divide o agrupa")
        notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        if not notes:
            add("INFO", si, "Sin notas del orador")
    if len(fonts) > 3:
        add("AVISO", 0, f"{len(fonts)} familias tipográficas en uso: {', '.join(sorted(fonts))}")
    return {"file": path, "slides": len(prs.slides), "fonts": sorted(fonts), "issues": issues}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--strict", action="store_true")
    ap.add_argument("--no-info", action="store_true", help="Oculta mensajes INFO")
    ap.add_argument("--documento", action="store_true",
                    help="Deck para leer (oferta, informe): admite más texto por diapositiva")
    ap.add_argument("--max-words", type=int)
    a = ap.parse_args()
    mw = a.max_words or (170 if a.documento else 90)
    rep = check(a.pptx, max_words=mw, max_bullets=9 if a.documento else 7)
    if a.no_info:
        rep["issues"] = [i for i in rep["issues"] if i["severity"] != "INFO"]
    if a.json:
        print(json.dumps(rep, ensure_ascii=False, indent=2))
    else:
        print(f"QA de {rep['file']} — {rep['slides']} diapositivas — fuentes: "
              f"{', '.join(rep['fonts']) or '-'}")
        order = {"ERROR": 0, "AVISO": 1, "INFO": 2}
        info_slides = [i["slide"] for i in rep["issues"] if i["severity"] == "INFO"]
        for it in sorted(rep["issues"], key=lambda i: (order[i["severity"]], i["slide"])):
            if it["severity"] == "INFO":
                continue
            where = f"diap. {it['slide']}" if it["slide"] else "general"
            print(f"  [{it['severity']}] {where}: {it['message']}")
        if info_slides:
            print(f"  [INFO] Sin notas del orador en: {', '.join(map(str, info_slides))}")
        n_err = sum(1 for i in rep["issues"] if i["severity"] == "ERROR")
        n_av = sum(1 for i in rep["issues"] if i["severity"] == "AVISO")
        print(f"Resultado: {n_err} errores, {n_av} avisos")
    if a.strict and any(i["severity"] == "ERROR" for i in rep["issues"]):
        sys.exit(1)


if __name__ == "__main__":
    main()
